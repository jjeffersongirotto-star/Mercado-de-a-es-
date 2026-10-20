"""Planilha .xlsx do Grupo simulação (⬇ do widget do app, /api/widget/simgrupo.xlsx): aba Resumo com todos os modelos lado a lado e
uma aba por modelo (resumo, Posições atuais e todas as movimentações). openpyxl em modo write_only (streaming): pouca memória mesmo
com dezenas de milhares de linhas. Formatos: R$ (BRL), %, inteiros; títulos com fundo, linha de títulos congelada, larguras fixas."""
import io
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
PCT = '0.00%;[Red]-0.00%'
INT = '#,##0'
PRC = '"R$" #,##0.00'
HEAD_FILL, HEAD_FONT = PatternFill("solid", fgColor="0F172A"), Font(bold=True, color="FFFFFF")
TITLE_FONT, SUB_FONT, BOLD = Font(bold=True, size=14), Font(italic=True, color="64748B"), Font(bold=True)
TIPO = {"compra": "Compra", "venda": "Venda", "aporte": "Aporte", "provento": "Provento", "cdi_ap": "Aplicação", "cdi_res": "Resgate",
        "retirada": "Retirada", "imposto": "Imposto"}
RESUMO = [("Modelo", 34, None), ("Carteira", 10, INT), ("Valor aportado", 16, BRL), ("Valor atual", 16, BRL), ("Variação do dia (R$)", 18, BRL),
          ("Variação do dia (%)", 16, PCT), ("Rendimento (%)", 15, PCT)]
POS = [("Empresa", 28, None), ("Ticker", 10, None), ("Ações", 10, INT), ("Preço médio", 14, PRC), ("Investido", 15, BRL), ("Atual", 15, BRL), ("Resultado (%)", 14, PCT)]
MOV = [("Data", 12, "DD/MM/YYYY"), ("Tipo", 12, None), ("Origem/destino", 15, None), ("Valor ação", 13, PRC), ("Qtd. de ações", 13, INT), ("Valor", 15, BRL), ("Detalhe", 36, None)]


def _c(ws, v, fmt=None, font=None, fill=None, al=None):
    c = WriteOnlyCell(ws, value=v)
    if fmt: c.number_format = fmt
    if font: c.font = font
    if fill: c.fill = fill
    if al: c.alignment = al
    return c


def _head(ws, cols):
    return [_c(ws, h, font=HEAD_FONT, fill=HEAD_FILL, al=Alignment(horizontal="center", vertical="center", wrap_text=True)) for h, _, _ in cols]


def _row(ws, cols, vals):
    return [_c(ws, v, fmt=f if v is not None and not isinstance(v, str) else None) for (_, _, f), v in zip(cols, vals)]


def _widths(ws, cols):
    from openpyxl.utils import get_column_letter
    for i, (_, w, _) in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = w


def _date(s):
    from datetime import date
    try: y, m, d = map(int, str(s)[:10].split("-")); return date(y, m, d)
    except Exception: return s


def _sheet_name(n, used):
    b = "".join(ch for ch in n if ch not in '[]:*?/\\')[:31] or "Modelo"; k, x = 2, b
    while x.lower() in used: x = f"{b[:28]} {k}"; k += 1
    used.add(x.lower()); return x


def pct(a, b): return (a / b - 1) if b else None


def build(name, w, sg, hist):
    """name: nome do grupo; w: dados do widget (server._simw: valores com o preço de agora); sg: modelos.simgrupo; hist(m) -> movimentações."""
    wb = Workbook(write_only=True); used = {"resumo"}
    per = f"Desde {w['d0'][8:10]}/{w['d0'][5:7]}/{w['d0'][:4]} · dados até {w['d1'][8:10]}/{w['d1'][5:7]}/{w['d1'][:4]} · preços de {w.get('u') or '—'}"
    ws = wb.create_sheet("Resumo"); _widths(ws, RESUMO); ws.freeze_panes = "A4"
    ws.append([_c(ws, name, font=TITLE_FONT)]); ws.append([_c(ws, per, font=SUB_FONT)]); ws.append(_head(ws, RESUMO))
    for x in w["r"]:
        ws.append(_row(ws, RESUMO, [x["n"], x["np"], x["ap"], x["v"], x["dv"], x["dp"] / 100 if x["dp"] is not None else None, pct(x["v"], x["ap"])]))
    for x in w["r"]:
        ws = wb.create_sheet(_sheet_name(x["n"], used))
        _widths(ws, [(None, max(a[1], b[1]), None) for a, b in zip(POS, MOV)])  # write_only: larguras e congelamento antes da 1ª linha
        summ = [("Valor aportado", x["ap"], BRL), ("Valor atual", x["v"], BRL), ("Rendimento", pct(x["v"], x["ap"]), PCT), ("Variação do dia", x["dv"], BRL),
                ("Carteira", f"{x['np']} empresa(s)", None), ("Caixa no CDB", x["cash"], BRL), ("Compras / vendas", f"{x.get('c') or 0} / {x.get('s') or 0}", None),
                ("Imposto pago", x.get("ir") or 0, BRL)]
        ws.freeze_panes = "A" + str(len(summ) + 6)  # resumo e títulos das posições sempre à vista
        ws.append([_c(ws, f"{name} · {x['n']}", font=TITLE_FONT)]); ws.append([_c(ws, per, font=SUB_FONT)])
        for k, v, f in summ: ws.append([_c(ws, k, font=BOLD), _c(ws, v, fmt=f if not isinstance(v, str) else None)])
        ws.append([]); ws.append([_c(ws, "Posições atuais", font=Font(bold=True, size=12))]); ws.append(_head(ws, POS))
        for y in x["pos"]:
            inv = y["inv"]; pm = y.get("pm") or (inv / y["q"] if y["q"] else None)
            ws.append(_row(ws, POS, [y.get("s") or y["t"], y["t"], y["q"], pm, inv, y["v"], pct(y["v"], inv)]))
        if x["cash"] >= 0.01: ws.append(_row(ws, POS, ["Caixa no CDB", "", None, None, None, x["cash"], None]))
        ws.append([]); ws.append([_c(ws, "Movimentações (todas, das mais recentes para as mais antigas)", font=Font(bold=True, size=12))]); ws.append(_head(ws, MOV))
        for e in reversed(hist(x["m"])):
            k = e.get("k"); cdi = k in ("cdi_ap", "cdi_res")
            ws.append(_row(ws, MOV, [_date(e.get("d")), TIPO.get(k, k), e.get("t") or ("CDB" if cdi else "—"),
                                     e.get("p") if e.get("p") is not None and k not in ("cdi_ap", "cdi_res", "imposto", "aporte") else None,
                                     e.get("q"), e.get("v"), e.get("x") or ""]))
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()
